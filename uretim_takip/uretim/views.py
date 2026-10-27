from datetime import timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from openpyxl import Workbook
from openpyxl.styles import Font

from . import raporlar, servisler, uyarilar
from .forms import IslemBaslatForm, IslemBitirForm, MalzemeKullanimFormSet, TarihAraligiForm
from .models import Asama, IslemKaydi, Malzeme, OturumKaydi, Siparis, UretimLotu


def geri_don(request, varsayilan):
    """İşlem istasyon ekranından açıldıysa (?geri=/istasyon/3/) oraya döner."""
    geri = request.GET.get("geri", "")
    if geri and url_has_allowed_host_and_scheme(geri, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        return redirect(geri)
    return redirect(varsayilan)


def sadece_ofis(gorunum):
    """Raporlar gibi ofis ekranları: personel durumlu hesaplar ve yönetici. Tablet hesabı göremez."""
    @login_required
    def sarmal(request, *args, **kwargs):
        if not request.user.is_staff:
            raise PermissionDenied
        return gorunum(request, *args, **kwargs)
    return sarmal


@login_required
def pano(request):
    bugun = timezone.localdate()
    aktif_lotlar = UretimLotu.objects.exclude(durum__in=["tamamlandi", "iptal"])
    asamalar = Asama.objects.filter(aktif=True).annotate(
        islemde=Count("uretimlotu", filter=Q(uretimlotu__durum="islemde")),
        bekleyen=Count("uretimlotu", filter=Q(uretimlotu__durum="bekliyor")),
    ).order_by("sira")
    baglam = {
        "asamalar": asamalar,
        "yeni_lot_sayisi": aktif_lotlar.filter(son_asama__isnull=True).count(),
        "acik_islemler": IslemKaydi.objects.filter(bitis__isnull=True).select_related("lot", "asama", "personel"),
        "bugun_personel": raporlar.personel_ozeti(bugun, bugun),
        "geciken_siparisler": Siparis.objects.filter(
            durum__in=["acik", "uretimde"], teslim_tarihi__lt=bugun
        ).select_related("musteri"),
        "yaklasan_siparisler": Siparis.objects.filter(
            durum__in=["acik", "uretimde"], teslim_tarihi__gte=bugun, teslim_tarihi__lte=bugun + timedelta(days=7)
        ).select_related("musteri"),
        "gecikmeler": uyarilar.gecikmeler(),
        "azalan_stoklar": uyarilar.azalan_stoklar() if request.user.is_staff else [],
    }
    return render(request, "uretim/pano.html", baglam)


@login_required
def islem_baslat(request):
    baslangic = {}
    if "lot" in request.GET:
        lot = UretimLotu.objects.filter(pk=request.GET["lot"]).first()
        if lot:
            sonrakiler = lot.siradaki_asamalar
            # Alternatif istasyonlar varsa (Biye / Lanfel) çalışan kendisi seçer.
            baslangic = {"lot": lot, "asama": sonrakiler[0] if len(sonrakiler) == 1 else None}

    form = IslemBaslatForm(request.POST or None, initial=baslangic)
    if request.method == "POST" and form.is_valid():
        v = form.cleaned_data
        try:
            islem = servisler.islem_baslat(
                lot=v["lot"], asama=v["asama"], personel=v["personel"], aciklama=v["aciklama"],
                makine=v.get("makine"),
            )
        except ValidationError as hata:
            form.add_error(None, hata)
        else:
            messages.success(request, f"{islem.personel} – {islem.lot} işinde “{islem.asama.ad}” başladı.")
            return geri_don(request, "pano")
    return render(request, "uretim/islem_baslat.html", {"form": form})


@login_required
def islem_bitir(request, pk):
    islem = get_object_or_404(IslemKaydi.objects.select_related("lot", "asama", "personel"), pk=pk)
    if islem.tamamlandi:
        messages.info(request, "Bu işlem zaten bitirilmiş.")
        return redirect("lot_detay", pk=islem.lot_id)

    form = IslemBitirForm(
        request.POST or None,
        initial={"personel": islem.personel, "cikis_kg": islem.giris_kg, "cikis_m2": islem.giris_m2},
    )
    # Malzeme bölümü gönderilmediyse (eski açık sayfa) malzeme kullanılmadı sayılır.
    malzeme_verisi = request.POST if "malzeme-TOTAL_FORMS" in request.POST else None
    malzemeler = MalzemeKullanimFormSet(malzeme_verisi, prefix="malzeme")
    if request.method == "POST" and form.is_valid() and (malzeme_verisi is None or malzemeler.is_valid()):
        v = form.cleaned_data
        kullanilan = [(m.cleaned_data["malzeme"], m.cleaned_data["miktar"])
                      for m in malzemeler if malzeme_verisi is not None and m.cleaned_data.get("malzeme")]
        try:
            servisler.islem_bitir(
                islem=islem, personel=v["personel"], cikis_kg=v["cikis_kg"], cikis_m2=v["cikis_m2"],
                aciklama=v["aciklama"], malzemeler=kullanilan,
            )
        except ValidationError as hata:
            form.add_error(None, hata)
        else:
            messages.success(request, f"{islem.lot} – “{islem.asama.ad}” bitti.")
            return geri_don(request, "pano")
    malzeme_kodlari = dict(Malzeme.objects.filter(aktif=True).values_list("kod", "pk"))
    return render(request, "uretim/islem_bitir.html",
                  {"form": form, "islem": islem, "malzemeler": malzemeler, "malzeme_kodlari": malzeme_kodlari})


@login_required
def lot_listesi(request):
    durum = request.GET.get("durum", "aktif")
    lotlar = UretimLotu.objects.select_related("siparis__musteri", "son_asama").prefetch_related("kaynaklar__ham_deri__deri_turu")
    if durum == "aktif":
        lotlar = lotlar.exclude(durum__in=["tamamlandi", "iptal"])
    elif durum in dict(UretimLotu.DURUM):
        lotlar = lotlar.filter(durum=durum)
    ara = request.GET.get("ara", "").strip()
    if ara:
        lotlar = lotlar.filter(
            Q(lot_no__icontains=ara) | Q(kaynaklar__ham_deri__lot_no__icontains=ara) | Q(siparis__siparis_no__icontains=ara)
        )
    return render(
        request, "uretim/lot_listesi.html",
        {"lotlar": lotlar.distinct()[:300], "durum": durum, "ara": ara, "durumlar": UretimLotu.DURUM},
    )


@login_required
def lot_detay(request, pk):
    lot = get_object_or_404(
        UretimLotu.objects.select_related("siparis__musteri"), pk=pk
    )
    islemler = raporlar.bekleme_suresi_ekle(
        list(lot.islemler.select_related("asama", "makine", "personel", "bitiren").order_by("baslangic"))
    )
    for islem in islemler:
        islem.bekleme_metni = uyarilar.sure_metni(islem.bekleme) if islem.bekleme is not None else ""
    return render(request, "uretim/lot_detay.html", {"lot": lot, "islemler": islemler})


def _tarih_araligi(request):
    bugun = timezone.localdate()
    form = TarihAraligiForm(request.GET or {"baslangic": bugun.replace(day=1), "bitis": bugun})
    if form.is_valid():
        return form, form.cleaned_data["baslangic"], form.cleaned_data["bitis"]
    return form, bugun.replace(day=1), bugun


@sadece_ofis
def rapor(request):
    form, bas, son = _tarih_araligi(request)
    personel = raporlar.personel_ozeti(bas, son)
    asama = raporlar.asama_ozeti(bas, son)
    makine = raporlar.makine_ozeti(bas, son)
    personel_istasyon = raporlar.personel_istasyon_ozeti(bas, son)
    for satir in personel + asama + makine + personel_istasyon:
        satir["sure_metni"] = raporlar.sure_metni(satir["toplam_sure"])
    bekleme = raporlar.bekleme_ozeti(bas, son)
    for satir in bekleme:
        satir["ortalama_metni"] = uyarilar.sure_metni(satir["ortalama"])
        satir["en_uzun_metni"] = uyarilar.sure_metni(satir["en_uzun"])
    return render(
        request, "uretim/rapor.html",
        {"form": form, "personel": personel, "asama": asama, "makine": makine, "bas": bas, "son": son,
         "personel_istasyon": personel_istasyon, "bekleme": bekleme},
    )


@sadece_ofis
def rapor_excel(request):
    _, bas, son = _tarih_araligi(request)
    kitap = Workbook()
    kalin = Font(bold=True)

    def sayfa_yaz(sayfa, basliklar, satirlar):
        sayfa.append(basliklar)
        for hucre in sayfa[1]:
            hucre.font = kalin
        for satir in satirlar:
            sayfa.append(satir)
        for sutun in sayfa.columns:
            sayfa.column_dimensions[sutun[0].column_letter].width = max(12, len(str(sutun[0].value)) + 4)

    ozet_baslik = ["İşlem", "Giriş kg", "Çıkış kg", "Fire kg", "Giriş m²", "Çıkış m²", "Fire m²", "Fire % (m²)", "Süre"]

    def ozet_satiri(s):
        return [
            s["islem_sayisi"], s["giris_kg"], s["cikis_kg"], s["fire_kg"], s["giris_m2"], s["cikis_m2"],
            s["fire_m2"], s["fire_yuzde_m2"], raporlar.sure_metni(s["toplam_sure"]),
        ]

    sayfa = kitap.active
    sayfa.title = "Personel"
    sayfa_yaz(sayfa, ["Personel"] + ozet_baslik,
              [[s["personel__ad_soyad"]] + ozet_satiri(s) for s in raporlar.personel_ozeti(bas, son)])

    sayfa_yaz(kitap.create_sheet("Aşamalar"), ["Aşama"] + ozet_baslik,
              [[s["asama__ad"]] + ozet_satiri(s) for s in raporlar.asama_ozeti(bas, son)])

    sayfa_yaz(kitap.create_sheet("Makineler"), ["Makine"] + ozet_baslik,
              [[s["makine__ad"]] + ozet_satiri(s) for s in raporlar.makine_ozeti(bas, son)])
    sayfa_yaz(kitap.create_sheet("Personel-İstasyon"), ["Personel", "İstasyon"] + ozet_baslik,
              [[s["personel__ad_soyad"], s["asama__ad"]] + ozet_satiri(s)
               for s in raporlar.personel_istasyon_ozeti(bas, son)])
    sayfa_yaz(kitap.create_sheet("İstasyonlar arası bekleme"), ["Geçiş", "İş sayısı", "Ortalama bekleme", "En uzun"],
              [[s["gecis"], s["sayi"], uyarilar.sure_metni(s["ortalama"]), uyarilar.sure_metni(s["en_uzun"])]
               for s in raporlar.bekleme_ozeti(bas, son)])

    detay = raporlar.tamamlanan_islemler(bas, son).select_related("lot", "asama", "makine", "personel", "bitiren")
    yerel = timezone.localtime
    sayfa_yaz(
        kitap.create_sheet("İşlem detayı"),
        ["İş", "Aşama", "Makine", "Personel", "Bitiren", "Başlangıç", "Bitiş", "Giriş kg", "Çıkış kg", "Giriş m²", "Çıkış m²", "Not"],
        [
            [i.lot.lot_no, i.asama.ad, i.makine.ad if i.makine else "", i.personel.ad_soyad, i.bitiren.ad_soyad if i.bitiren else "",
             yerel(i.baslangic).replace(tzinfo=None), yerel(i.bitis).replace(tzinfo=None),
             i.giris_kg, i.cikis_kg, i.giris_m2, i.cikis_m2, i.aciklama]
            for i in detay.order_by("bitis")
        ],
    )

    yanit = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    yanit["Content-Disposition"] = f'attachment; filename="uretim_raporu_{bas}_{son}.xlsx"'
    kitap.save(yanit)
    return yanit


@sadece_ofis
def tablet_girisleri(request):
    """Hangi hesap ne zaman giriş yaptı / hangi istasyon ekranını açtı."""
    bugun = timezone.localdate()
    form = TarihAraligiForm(request.GET or {"baslangic": bugun, "bitis": bugun})
    bas, son = (form.cleaned_data["baslangic"], form.cleaned_data["bitis"]) if form.is_valid() else (bugun, bugun)
    alt, ust = raporlar._aralik(bas, son)
    kayitlar = OturumKaydi.objects.filter(zaman__gte=alt, zaman__lt=ust).select_related("asama")
    hesap = request.GET.get("hesap", "")
    if hesap:
        kayitlar = kayitlar.filter(kullanici_adi=hesap)
    hesaplar = OturumKaydi.objects.order_by("kullanici_adi").values_list("kullanici_adi", flat=True).distinct()

    ozet = {}
    for k in kayitlar.order_by("zaman"):
        o = ozet.setdefault(k.kullanici_adi, {"hesap": k.kullanici_adi, "ilk": k.zaman, "son": k.zaman,
                                               "istasyonlar": [], "cihaz": k.cihaz})
        o["son"] = k.zaman
        if k.asama and k.asama.ad not in o["istasyonlar"]:
            o["istasyonlar"].append(k.asama.ad)
    return render(request, "uretim/tablet_girisleri.html", {
        "form": form, "kayitlar": kayitlar[:500], "ozet": list(ozet.values()), "hesap": hesap, "hesaplar": hesaplar,
    })
