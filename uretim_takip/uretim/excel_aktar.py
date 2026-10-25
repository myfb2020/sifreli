"""Excel'den toplu veri aktarımı.

Şablondaki her sayfa bir kayıt türüdür. Aktarım tek bir veritabanı işlemi içinde yapılır:
önizlemede (kaydet=False) her şey denenir ve sonra geri alınır. Böylece önizleme, gerçek
kayıtla birebir aynı sonucu ve hataları gösterir. Kaydederken tek bir hata bile varsa hiçbir
şey kaydedilmez.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from .models import (
    Asama, DeriTuru, HamDeriGirisi, LotKaynagi, Musteri, Personel, Siparis, Tedarikci, UretimLotu, pin_dogrulayici,
)


class SatirHatasi(Exception):
    pass


# --- Değer dönüştürme ---------------------------------------------------------

_TR = str.maketrans("İIıŞşĞğÜüÖöÇç²", "iiissgguuoocc2")


def normal(metin):
    """Başlık ve sayfa adlarını karşılaştırmak için: 'Başlangıç m²' -> 'baslangicm2'."""
    return re.sub(r"[^a-z0-9]", "", str(metin or "").translate(_TR).lower())


def metin(deger):
    if deger is None:
        return ""
    if isinstance(deger, float) and deger.is_integer():
        deger = int(deger)
    return str(deger).strip()


def sinirli(deger, alan, uzunluk):
    s = metin(deger)
    if len(s) > uzunluk:
        raise SatirHatasi(f"“{alan}” en fazla {uzunluk} karakter olabilir")
    return s


def kod(deger, alan):
    return sinirli(deger, alan, 20).upper()


def sayi(deger, alan):
    if deger is None or deger == "":
        return None
    if isinstance(deger, (int, float, Decimal)) and not isinstance(deger, bool):
        sonuc = Decimal(str(deger))
    else:
        s = str(deger).strip().replace(" ", "")
        if "," in s and "." in s:  # 1.234,5 (Türkçe) ya da 1,234.5 (İngilizce)
            s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
        else:
            s = s.replace(",", ".")
        try:
            sonuc = Decimal(s)
        except InvalidOperation:
            raise SatirHatasi(f"“{alan}” sayı olmalı (yazılan: {deger})")
    if sonuc < 0:
        raise SatirHatasi(f"“{alan}” negatif olamaz")
    return sonuc.quantize(Decimal("0.01"))


def tam_sayi(deger, alan):
    s = sayi(deger, alan)
    if s is None:
        return None
    if s != s.to_integral_value():
        raise SatirHatasi(f"“{alan}” tam sayı olmalı (yazılan: {deger})")
    return int(s)


def tarih(deger, alan):
    if deger in (None, ""):
        return None
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    s = str(deger).strip()
    for bicim in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(s, bicim).date()
        except ValueError:
            pass
    raise SatirHatasi(f"“{alan}” tarih olmalı, ör. 25.10.2026 (yazılan: {deger})")


def secim(deger, alan, secenekler):
    """Görünen ada ya da koda göre seçim: 'Sevk edildi' / 'sevk' -> 'sevk'."""
    if deger in (None, ""):
        return None
    aranan = normal(deger)
    for kod, ad in secenekler:
        if aranan in (normal(kod), normal(ad)):
            return kod
    gecerli = ", ".join(ad for _, ad in secenekler)
    raise SatirHatasi(f"“{alan}” şunlardan biri olmalı: {gecerli} (yazılan: {deger})")


# --- Sayfa tanımları ----------------------------------------------------------


@dataclass
class Sutun:
    baslik: str
    zorunlu: bool = False
    aciklama: str = ""
    ornek: str = ""
    takma_adlar: tuple = ()  # eski şablonlardaki başlıklar


@dataclass
class Sayfa:
    ad: str
    aciklama: str
    sutunlar: list
    isleyici: object = None
    takma_adlar: tuple = ()


@dataclass
class Rapor:
    eklenen: dict = field(default_factory=dict)
    guncellenen: dict = field(default_factory=dict)
    hatalar: list = field(default_factory=list)
    uyarilar: list = field(default_factory=list)
    yeni_tanimlar: list = field(default_factory=list)
    bulunan_sayfalar: list = field(default_factory=list)
    kaydedildi: bool = False

    @property
    def toplam(self):
        return sum(self.eklenen.values()) + sum(self.guncellenen.values())

    @property
    def ozet(self):
        adlar = list(dict.fromkeys([*self.eklenen, *self.guncellenen]))
        return [(ad, self.eklenen.get(ad, 0), self.guncellenen.get(ad, 0)) for ad in adlar]

    def say(self, sayfa, yeni):
        hedef = self.eklenen if yeni else self.guncellenen
        hedef[sayfa] = hedef.get(sayfa, 0) + 1


class Aktarici:
    def __init__(self):
        self.rapor = Rapor()

    # Ortak yardımcılar
    def _tanim(self, model, ad, tur):
        ad = sinirli(ad, tur, 60 if model is DeriTuru else 120)
        if not ad:
            return None
        nesne = model.objects.filter(ad__iexact=ad).first()
        if nesne is None:
            nesne = model.objects.create(ad=ad)
            self.rapor.yeni_tanimlar.append(f"{tur}: {ad}")
        return nesne

    @staticmethod
    def _zorunlu(satir, *alanlar):
        eksik = [a for a in alanlar if satir.get(a) in (None, "")]
        if eksik:
            raise SatirHatasi("Boş bırakılamaz: " + ", ".join(eksik))

    # Sayfa işleyicileri: her biri (yeni_mi) döner.
    def personel(self, s):
        self._zorunlu(s, "Ad Soyad")
        ad = sinirli(s["Ad Soyad"], "Ad Soyad", 100)
        pin = metin(s.get("PIN"))
        if pin:
            try:
                pin_dogrulayici(pin)
            except ValidationError:
                ipucu = " (Excel baştaki sıfırı silmiş olabilir; PIN sütununu Metin biçimine çevirin)" \
                    if isinstance(s.get("PIN"), (int, float)) and len(pin) < 4 else ""
                raise SatirHatasi("PIN 4-6 haneli rakam olmalı" + ipucu)
        p = Personel.objects.filter(ad_soyad__iexact=ad).first()
        yeni = p is None
        if yeni:
            if not pin:
                raise SatirHatasi("Yeni personel için PIN zorunlu")
            p = Personel(ad_soyad=ad)
        if pin:
            p.pin_ayarla(pin)
        if s.get("Görev") not in (None, ""):
            p.gorev = sinirli(s["Görev"], "Görev", 60)
        aktif = secim(s.get("Aktif"), "Aktif", [("evet", "Evet"), ("hayir", "Hayır")])
        if aktif:
            p.aktif = aktif == "evet"
        p.save()
        return yeni

    def _basit(self, model, s):
        self._zorunlu(s, "Ad")
        ad = sinirli(s["Ad"], "Ad", 120)
        nesne = model.objects.filter(ad__iexact=ad).first()
        yeni = nesne is None
        nesne = nesne or model(ad=ad)
        for baslik, alan in (("Telefon", "telefon"), ("Not", "notlar")):
            if s.get(baslik) not in (None, ""):
                setattr(nesne, alan, sinirli(s[baslik], baslik, 30) if alan == "telefon" else metin(s[baslik]))
        nesne.save()
        return yeni

    def tedarikci(self, s):
        return self._basit(Tedarikci, s)

    def musteri(self, s):
        return self._basit(Musteri, s)

    def ham_deri(self, s):
        self._zorunlu(s, "Tedarikçi", "Deri Türü", "Kg", "m²")
        lot_no = kod(s.get("Lot No"), "Lot No")
        h = HamDeriGirisi.objects.filter(lot_no=lot_no).first() if lot_no else None
        yeni = h is None
        h = h or HamDeriGirisi(lot_no=lot_no)
        h.tedarikci = self._tanim(Tedarikci, s["Tedarikçi"], "Tedarikçi")
        h.deri_turu = self._tanim(DeriTuru, s["Deri Türü"], "Deri türü")
        h.kg = sayi(s["Kg"], "Kg")
        h.m2 = sayi(s["m²"], "m²")
        h.tarih = tarih(s.get("Tarih"), "Tarih") or h.tarih
        h.adet = tam_sayi(s.get("Adet"), "Adet")
        h.kalite = secim(s.get("Kalite"), "Kalite", HamDeriGirisi.KALITE) or h.kalite
        h.notlar = metin(s.get("Not")) or h.notlar
        h.save()
        return yeni

    def siparis(self, s):
        self._zorunlu(s, "Müşteri", "Ürün", "Teslim Tarihi")
        no = kod(s.get("Sipariş No"), "Sipariş No")
        sp = Siparis.objects.filter(siparis_no=no).first() if no else None
        yeni = sp is None
        sp = sp or Siparis(siparis_no=no)
        sp.musteri = self._tanim(Musteri, s["Müşteri"], "Müşteri")
        sp.urun = sinirli(s["Ürün"], "Ürün", 200)
        sp.teslim_tarihi = tarih(s["Teslim Tarihi"], "Teslim Tarihi")
        sp.tarih = tarih(s.get("Tarih"), "Tarih") or sp.tarih
        sp.deri_turu = self._tanim(DeriTuru, s.get("Deri Türü"), "Deri türü") or sp.deri_turu
        sp.adet = tam_sayi(s.get("Adet"), "Adet") or sp.adet
        sp.toplam_m2 = sayi(s.get("Toplam m²"), "Toplam m²")
        sp.durum = secim(s.get("Durum"), "Durum", Siparis.DURUM) or sp.durum
        sp.notlar = metin(s.get("Not")) or sp.notlar
        sp.save()
        return yeni

    def lot(self, s, satir_no):
        self._zorunlu(s, "Stok Lot No", "Başlangıç kg", "Başlangıç m²")
        ham_no = kod(s["Stok Lot No"], "Stok Lot No")
        ham = HamDeriGirisi.objects.filter(lot_no=ham_no).first()
        if ham is None:
            raise SatirHatasi(f"“{ham_no}” numaralı stok lotu bulunamadı (Stok Lotları sayfasında veya sistemde olmalı)")
        siparis = None
        if metin(s.get("Sipariş No")):
            sp_no = kod(s["Sipariş No"], "Sipariş No")
            siparis = Siparis.objects.filter(siparis_no=sp_no).first()
            if siparis is None:
                raise SatirHatasi(f"“{sp_no}” numaralı sipariş bulunamadı")

        asama = None
        asama_degeri = metin(s.get("Son Tamamlanan Aşama"))
        if asama_degeri:
            aranan = normal(asama_degeri)
            for a in Asama.objects.all():
                if aranan in (normal(a.ad), str(a.sira)) or normal(a.ad).startswith(aranan):
                    asama = a
                    break
            if asama is None:
                raise SatirHatasi(f"“{asama_degeri}” adında bir aşama yok")

        durumlar = [kod for kod in UretimLotu.DURUM if kod[0] != "islemde"]
        durum = secim(s.get("Durum"), "Durum", durumlar)
        if durum is None:
            durum = "tamamlandi" if asama and asama.son_asama_mi else "bekliyor"

        lot_no = kod(s.get("İş No"), "İş No")
        lot = UretimLotu.objects.filter(lot_no=lot_no).first() if lot_no else None
        yeni = lot is None
        if not yeni and lot.islemler.filter(bitis__isnull=True).exists():
            raise SatirHatasi(f"{lot_no} işinde sistemde açık bir işlem var, Excel'den güncellenemez")
        lot = lot or UretimLotu(lot_no=lot_no)
        lot.siparis = siparis
        lot.baslangic_kg = sayi(s["Başlangıç kg"], "Başlangıç kg")
        lot.baslangic_m2 = sayi(s["Başlangıç m²"], "Başlangıç m²")
        lot.mevcut_kg = sayi(s.get("Mevcut kg"), "Mevcut kg")
        lot.mevcut_m2 = sayi(s.get("Mevcut m²"), "Mevcut m²")
        if lot.mevcut_kg is None:
            lot.mevcut_kg = lot.baslangic_kg
        if lot.mevcut_m2 is None:
            lot.mevcut_m2 = lot.baslangic_m2
        lot.son_asama = asama
        lot.durum = durum
        lot.notlar = metin(s.get("Not")) or lot.notlar
        lot.save()
        lot.kaynaklar.all().delete()
        LotKaynagi.objects.create(uretim_lotu=lot, ham_deri=ham, kg=lot.baslangic_kg, m2=lot.baslangic_m2)

        if ham.kalan_kg < 0 or ham.kalan_m2 < 0:
            self.rapor.uyarilar.append(
                f"Devam Eden İşler, satır {satir_no}: {ham.lot_no} stok lotundan girişten fazla malzeme alınmış "
                f"(kalan {ham.kalan_kg} kg / {ham.kalan_m2} m²)."
            )
        return yeni


SAYFALAR = [
    Sayfa("Personel", "Çalışanlar. Ad soyada göre eşleşir: aynı isim varsa güncellenir.", [
        Sutun("Ad Soyad", True, "Çalışanın adı soyadı", "Ahmet Yılmaz"),
        Sutun("PIN", False, "4-6 haneli rakam. Yeni personel için zorunlu, mevcut personelde boşsa değişmez.", "1234"),
        Sutun("Görev", False, "", "Kesim"),
        Sutun("Aktif", False, "Evet / Hayır (boşsa Evet)", "Evet"),
    ], "personel"),
    Sayfa("Tedarikçiler", "Ham deri aldığınız firmalar. Ada göre eşleşir.", [
        Sutun("Ad", True, "Firma adı", "Anadolu Deri Ltd."),
        Sutun("Telefon", False, "", "0532 000 00 00"),
        Sutun("Not", False),
    ], "tedarikci", ("tedarikci",)),
    Sayfa("Müşteriler", "Ada göre eşleşir.", [
        Sutun("Ad", True, "Müşteri adı", "Örnek Mobilya A.Ş."),
        Sutun("Telefon", False),
        Sutun("Not", False),
    ], "musteri", ("musteri",)),
    Sayfa("Stok Lotları", "Stoğa gelen ham deri partileri. Tedarikçi ve deri türü sistemde yoksa otomatik eklenir.", [
        Sutun("Lot No", False, "Boş bırakılırsa sistem numara verir (HD-2026-0001). Eski numaralarınızı yazarsanız onlar kullanılır.", "HD-2026-0001"),
        Sutun("Tarih", False, "Giriş tarihi (boşsa bugün)", "01.10.2026"),
        Sutun("Tedarikçi", True, "", "Anadolu Deri Ltd."),
        Sutun("Deri Türü", True, "Koyun, Kuzu, Keçi…", "Kuzu"),
        Sutun("Kalite", False, "A / B / C / Karışık", "A"),
        Sutun("Adet", False, "Deri adedi", "200"),
        Sutun("Kg", True, "Toplam ağırlık", "640"),
        Sutun("m²", True, "Toplam alan", "110,5"),
        Sutun("Not", False),
    ], "ham_deri", ("hamderigirisleri", "hamderi", "hammadde", "stok")),
    Sayfa("Siparişler", "Müşteri siparişleri. Müşteri sistemde yoksa otomatik eklenir.", [
        Sutun("Sipariş No", False, "Boş bırakılırsa sistem numara verir (SP-2026-0001)", "SP-2026-0001"),
        Sutun("Müşteri", True, "", "Örnek Mobilya A.Ş."),
        Sutun("Tarih", False, "Sipariş tarihi (boşsa bugün)", "05.10.2026"),
        Sutun("Ürün", True, "Ürün açıklaması", "200x300 krem kuzu patchwork halı"),
        Sutun("Deri Türü", False, "", "Kuzu"),
        Sutun("Adet", False, "", "4"),
        Sutun("Toplam m²", False, "", "24"),
        Sutun("Teslim Tarihi", True, "", "30.10.2026"),
        Sutun("Durum", False, "Açık / Üretimde / Tamamlandı / Sevk edildi / İptal (boşsa Açık)", "Açık"),
        Sutun("Not", False),
    ], "siparis", ("siparis",)),
    Sayfa("Devam Eden İşler", "Sistemi kurarken üretimde yarım kalan işler. Stok lotu ve sipariş numaraları yukarıdaki sayfalarda ya da sistemde bulunmalı.", [
        Sutun("İş No", False, "Boş bırakılırsa sistem numara verir (IS-2026-0001)", "IS-2026-0001", ("Lot No",)),
        Sutun("Stok Lot No", True, "Malzemenin alındığı stok lotu (Stok Lotları sayfasındaki Lot No)", "HD-2026-0001", ("Ham Deri Lot No",)),
        Sutun("Sipariş No", False, "", "SP-2026-0001"),
        Sutun("Başlangıç kg", True, "", "320"),
        Sutun("Başlangıç m²", True, "", "55"),
        Sutun("Mevcut kg", False, "Şu anki ağırlık (boşsa başlangıç)", "284"),
        Sutun("Mevcut m²", False, "Şu anki alan (boşsa başlangıç)", "50,5"),
        Sutun("Son Tamamlanan Aşama", False, "İşin bitirdiği son aşama: adı (ör. Tabaklama) ya da sıra no (ör. 30). Hiç başlamadıysa boş.", "Tabaklama"),
        Sutun("Durum", False, "Bekliyor / Tamamlandı / İptal (boşsa otomatik)", ""),
        Sutun("Not", False),
    ], "lot", ("uretimlotlari", "uretimlotu", "lotlar", "lot", "isler")),
]


def _sayfa_bul(kitap, sayfa):
    adlar = {normal(sayfa.ad), *sayfa.takma_adlar}
    for ws in kitap.worksheets:
        if normal(ws.title) in adlar:
            return ws
    return None


def _satirlar(ws, sayfa):
    """Başlık satırını bulur (ilk 10 satırda) ve her veri satırını {Başlık: değer} olarak verir."""
    beklenen = {normal(s.baslik): s.baslik for s in sayfa.sutunlar}
    beklenen.update({normal(t): s.baslik for s in sayfa.sutunlar for t in s.takma_adlar})
    for baslik_satiri, satir in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        eslesme = {i: beklenen[normal(h)] for i, h in enumerate(satir) if normal(h) in beklenen}
        if len(eslesme) >= max(1, len([s for s in sayfa.sutunlar if s.zorunlu])):
            break
    else:
        return None, []
    tanimsiz = [h for i, h in enumerate(satir) if h not in (None, "") and i not in eslesme]
    veriler = []
    for satir_no, degerler in enumerate(
        ws.iter_rows(min_row=baslik_satiri + 1, values_only=True), start=baslik_satiri + 1
    ):
        kayit = {eslesme[i]: d for i, d in enumerate(degerler) if i in eslesme}
        if any(metin(d) for d in kayit.values()):
            veriler.append((satir_no, kayit))
    return tanimsiz, veriler


def ice_aktar(dosya, kaydet=False):
    try:
        kitap = load_workbook(dosya, data_only=True, read_only=True)
    except Exception:
        rapor = Rapor()
        rapor.hatalar.append(("Dosya", "-", "Dosya okunamadı. Excel dosyasını .xlsx olarak kaydedip tekrar deneyin."))
        return rapor

    aktarici = Aktarici()
    rapor = aktarici.rapor
    with transaction.atomic():
        for sayfa in SAYFALAR:
            ws = _sayfa_bul(kitap, sayfa)
            if ws is None:
                continue
            tanimsiz, veriler = _satirlar(ws, sayfa)
            if tanimsiz is None:
                rapor.hatalar.append((sayfa.ad, 1, "Başlık satırı bulunamadı. Şablondaki başlıkları kullanın."))
                continue
            if veriler:
                rapor.bulunan_sayfalar.append(f"{sayfa.ad} ({len(veriler)} satır)")
            if tanimsiz:
                rapor.uyarilar.append(f"{sayfa.ad}: tanınmayan sütunlar atlandı: {', '.join(map(str, tanimsiz))}")
            isleyici = getattr(aktarici, sayfa.isleyici)
            for satir_no, kayit in veriler:
                try:
                    with transaction.atomic():
                        if sayfa.isleyici == "lot":
                            yeni = isleyici(kayit, satir_no)
                        else:
                            yeni = isleyici(kayit)
                    rapor.say(sayfa.ad, yeni)
                except SatirHatasi as hata:
                    rapor.hatalar.append((sayfa.ad, satir_no, str(hata)))
                except ValidationError as hata:
                    rapor.hatalar.append((sayfa.ad, satir_no, "; ".join(hata.messages)))
                except IntegrityError:
                    rapor.hatalar.append((sayfa.ad, satir_no, "Kayıt veritabanındaki başka bir kayıtla çakışıyor"))
        kitap.close()

        if not rapor.bulunan_sayfalar and not rapor.hatalar:
            rapor.hatalar.append(("Dosya", "-", "Aktarılacak veri bulunamadı. Sayfa adları şablondaki gibi olmalı "
                                                "(Personel, Stok Lotları, …) ve başlıkların altında veri bulunmalı."))
        if kaydet and not rapor.hatalar:
            rapor.kaydedildi = True
        else:
            transaction.set_rollback(True)
    return rapor


def sablon_olustur():
    kitap = Workbook()
    kalin, beyaz = Font(bold=True), Font(bold=True, color="FFFFFF")
    zorunlu_dolgu = PatternFill("solid", fgColor="7A4B2A")
    istege_dolgu = PatternFill("solid", fgColor="B08968")

    aciklama = kitap.active
    aciklama.title = "Açıklama"
    aciklama.append(["Üretim Takip – Excel aktarım şablonu"])
    aciklama["A1"].font = Font(bold=True, size=14)
    for satir in [
        [],
        ["1. Mevcut Excel'inizdeki verileri ilgili sayfalara kopyalayıp yapıştırın. Başlık satırını değiştirmeyin."],
        ["2. Koyu kahverengi başlıklar zorunlu, açık kahverengiler isteğe bağlıdır. Kullanmadığınız sayfaları boş bırakın."],
        ["3. Sisteme 'Excel'den aktar' sayfasından yükleyin. Önce önizleme gösterilir, onaylamadan hiçbir şey kaydedilmez."],
        ["4. Aynı dosyayı tekrar yüklerseniz kayıtlar çoğalmaz; Lot No / Sipariş No / ad eşleşen kayıtlar güncellenir."],
        ["5. Sayılarda 12,5 ya da 12.5 yazabilirsiniz. Tarihler: 25.10.2026"],
        [],
        ["Sayfa", "Sütun", "Zorunlu", "Açıklama", "Örnek"],
    ]:
        aciklama.append(satir)
    for hucre in aciklama[aciklama.max_row]:
        hucre.font = kalin
    for sayfa in SAYFALAR:
        aciklama.append([sayfa.ad, "", "", sayfa.aciklama])
        aciklama.cell(aciklama.max_row, 1).font = kalin
        for s in sayfa.sutunlar:
            aciklama.append(["", s.baslik, "Evet" if s.zorunlu else "", s.aciklama, s.ornek])
    for harf, genislik in zip("ABCDE", (16, 24, 9, 90, 34)):
        aciklama.column_dimensions[harf].width = genislik

    for sayfa in SAYFALAR:
        ws = kitap.create_sheet(sayfa.ad)
        ws.append([s.baslik for s in sayfa.sutunlar])
        for hucre, s in zip(ws[1], sayfa.sutunlar):
            hucre.font = beyaz
            hucre.fill = zorunlu_dolgu if s.zorunlu else istege_dolgu
            hucre.alignment = Alignment(horizontal="center")
            if s.aciklama or s.ornek:
                hucre.comment = Comment(f"{s.aciklama}\nÖrnek: {s.ornek}".strip(), "Üretim Takip")
            ws.column_dimensions[hucre.column_letter].width = max(14, len(s.baslik) + 6)
        ws.freeze_panes = "A2"
    return kitap
